# main.py
from docx import Document
import re
import os
import json
import argparse
from typing import List, Tuple

def extract_section_between_titles(docx_path, start_title, stop_keywords):
    doc = Document(docx_path)
    extracting = False
    result = ""

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        if not extracting and start_title.lower() in text.lower():
            extracting = True
            result += text + "\n"
            continue

        if extracting and any(keyword.lower() in text.lower() for keyword in stop_keywords):
            break

        if extracting:
            result += text + "\n"

    return result


def write_text_to_docx(text, output_path):
    document = Document()
    for line in text.strip().split("\n"):
        document.add_paragraph(line)
    document.save(output_path)


def extract_bilesen_and_use_cases(docx_path2, uc_limit=None, extra_stop_titles=None):
    doc = Document(docx_path2)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    bilesen_tasarimi = ""
    use_cases = []

    uc_pattern = re.compile(r"UC\s*\d*+\s*[-–]", re.IGNORECASE)

    in_bilesen = False
    in_use_case = False
    current_uc = None
    uc_collected = 0

    for line in paragraphs:
        line_lower = line.lower()

        if not in_bilesen:
            if "bileşen tasarımı" in line_lower:
                in_bilesen = True
            continue

        if in_bilesen and not in_use_case:
            if uc_pattern.match(line):
                in_use_case = True
                current_uc = {"title": line, "description": ""}
                uc_collected += 1
                continue
            else:
                bilesen_tasarimi += line + "\n"
                continue

        if in_use_case:
            if uc_pattern.match(line):
                if current_uc:
                    use_cases.append(current_uc)
                    current_uc = None
                if uc_limit and uc_collected >= uc_limit:
                    break
                current_uc = {"title": line, "description": ""}
                uc_collected += 1
            elif extra_stop_titles and any(stop.lower() in line_lower for stop in extra_stop_titles):
                if current_uc and (not uc_limit or uc_collected <= uc_limit):
                    use_cases.append(current_uc)
                    current_uc = None
                break
            else:
                if not uc_limit or uc_collected <= uc_limit:
                    current_uc["description"] += line + "\n"

    if current_uc and (not uc_limit or uc_collected <= uc_limit):
        use_cases.append(current_uc)

    return bilesen_tasarimi.strip(), use_cases


def parse_csv_like_list(v: str) -> List[str]:
    if not v:
        return []
    return [s.strip() for s in v.split(",") if s.strip()]


def build_arg_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="DOCX içinden 'Bileşen Tasarımı' ve UC'leri çıkarıp use_cases.json üretir."
    )
    p.add_argument("--doc", required=True, help="Kaynak DOCX dosya yolu")
    p.add_argument("--uc-limit", type=int, default=3, help="Maksimum alınacak Use-Case sayısı (varsayılan: 3)")
    p.add_argument("--start-title", default="Bileşen Tasarımı", help="Başlangıç başlığı (varsayılan: 'Bileşen Tasarımı')")
    p.add_argument(
        "--stop-titles",
        default="Veri Tasarımı,Sistem Mimarisi",
        help="Başlığı bitirecek üst başlıklar (virgülle ayır) (varsayılan: 'Veri Tasarımı,Sistem Mimarisi')"
    )
    p.add_argument(
        "--extra-stop-titles",
        default="hata dönüş kodları,genel tasarım modeli,döküman geliştirmeleri,sms geliştirmeleri,rule engine gereksinimleri,job listesi",
        help="UC toplamayı durduracak ek başlıklar (virgülle ayır)"
    )
    p.add_argument("--temp-docx", default="use_cases.docx", help="Geçici docx adı (varsayılan: use_cases.docx)")
    p.add_argument("--output-json", default="use_cases.json", help="Çıktı JSON adı (varsayılan: use_cases.json)")
    return p


def run(doc_path: str,
        uc_limit: int,
        start_title: str,
        stop_titles: List[str],
        extra_stop_titles: List[str],
        temp_docx: str,
        output_json: str) -> Tuple[str, list]:
    if not os.path.exists(doc_path):
        raise FileNotFoundError(f"Kaynak DOCX bulunamadı: {doc_path}")

    # 1) Belirli kısmı al
    section_text = extract_section_between_titles(doc_path, start_title, stop_titles)

    # 2) Geçici DOCX'e yaz
    write_text_to_docx(section_text, temp_docx)

    # 3) Use-Case ayrıştır
    bilesen_tasarimi, use_cases = extract_bilesen_and_use_cases(
        temp_docx,
        uc_limit=uc_limit,
        extra_stop_titles=extra_stop_titles
    )

    # 4) JSON kaydet
    with open(output_json, "w", encoding="utf-8") as f:
        json.dump({
            "bilesen_tasarimi": bilesen_tasarimi.strip(),
            "use_cases": use_cases
        }, f, ensure_ascii=False, indent=2)

    # 5) Geçici dosyayı sil
    if os.path.exists(temp_docx):
        os.remove(temp_docx)

    print(f"✅ {output_json} üretildi. UC sayısı: {len(use_cases)}")
    return bilesen_tasarimi, use_cases


if __name__ == "__main__":
    args = build_arg_parser().parse_args()
    stop_titles = parse_csv_like_list(args.stop_titles)
    extra_stop_titles = parse_csv_like_list(args.extra_stop_titles)

    run(
        doc_path=args.doc,
        uc_limit=args.uc_limit,
        start_title=args.start_title,
        stop_titles=stop_titles,
        extra_stop_titles=extra_stop_titles,
        temp_docx=args.temp_docx,
        output_json=args.output_json
    )
